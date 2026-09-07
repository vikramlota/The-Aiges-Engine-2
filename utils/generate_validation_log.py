"""
Turns a batch of pipeline results (from ig_pipeline.py's
run_integrated_pipeline / master_audit_results) into a workbook for the
Week 1-2 validation exercise: read each real post yourself, form your own
independent verdict, THEN compare it to what the engine said.

Usage from ig_pipeline.py, after a real run:

    from utils.generate_validation_log import write_validation_log
    write_validation_log(master_audit_results, "data/validation_log.xlsx")

Or standalone with no arguments to see a sample workbook built from
placeholder data (useful for previewing the layout before you have real
pipeline output).
"""
import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

NAVY = "1F3A5F"
GOLD = "B08D57"
SLATE = "44546A"
LIGHT = "F2F4F7"
WHITE = "FFFFFF"
GREEN = "C6EFCE"
RED = "FFC7CE"
AMBER = "FFEB9C"
FONT_NAME = "Calibri"
thin = Side(style="thin", color="D9DEE4")
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

VERDICT_OPTIONS = ["COMPLIANT", "FLAGGED", "NEEDS EXPERT REVIEW", "PENDING REVIEW"]
RISK_OPTIONS = ["CRITICAL", "HIGH", "MEDIUM", "LOW", "ADVISORY", "NONE"]
DISAGREEMENT_TYPES = [
    "N/A - Match",
    "False Positive - engine over-flagged",
    "False Negative - engine missed it",
    "Right verdict, wrong category or reason",
    "Risk level off",
    "Other",
]


def _title(ws, cell, text, size=16, color=NAVY):
    c = ws[cell]
    c.value = text
    c.font = Font(name=FONT_NAME, bold=True, size=size, color=color)


def _body(ws, cell, text, bold=False, italic=False, size=11, color="222222"):
    c = ws[cell]
    c.value = text
    c.font = Font(name=FONT_NAME, bold=bold, italic=italic, size=size, color=color)
    c.alignment = Alignment(wrap_text=True, vertical="top")


def _header_cell(ws, row, col, text, width, fill=NAVY):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(name=FONT_NAME, bold=True, size=10, color=WHITE)
    c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    c.border = border_all
    ws.column_dimensions[get_column_letter(col)].width = width


def write_validation_log(reports, filename="data/validation_log.xlsx"):
    """
    reports: a list of objects (or dicts) with these fields --
    post_id, influencer_handle, post_url, timestamp, caption_status,
    caption_flags, visual_status, risk_level, expert_review_flags.
    Matches ig_pipeline.py's UnifiedAuditReport exactly; dicts with the
    same keys work too.
    """
    out_path = Path(filename)
    if out_path.parent:
        out_path.parent.mkdir(parents=True, exist_ok=True)

    def get(r, key, default=""):
        if isinstance(r, dict):
            return r.get(key, default)
        return getattr(r, key, default)

    wb = Workbook()

    # ================= Sheet 1: How To Use =================
    ws1 = wb.active
    ws1.title = "How To Use"
    ws1.sheet_view.showGridLines = False
    ws1.column_dimensions["A"].width = 100

    _title(ws1, "A1", "Validation Log — Week 1-2 (Engine vs. Your Judgment)")
    _body(ws1, "A2", "Every row is one real post the pipeline already processed. Your job is to independently "
                     "decide what YOU think the verdict should be, before looking at what the engine said.",
          italic=True, color=SLATE)

    rows = [
        ("How to use this workbook", True, NAVY, 13),
        ("1. Open the \u201cValidation Log\u201d tab. Columns B\u2013E are already filled in from the real pipeline run.", False, "222222", 11),
        ("2. For each row: open the Post URL, read the actual post, and form your own opinion BEFORE touching columns I onward.", False, "222222", 11),
        ("3. Fill in \u201cYour Verdict\u201d (F) and, optionally, \u201cYour Risk Guess\u201d (G) and \u201cYour Reasoning\u201d (H) -- based only on your own judgment.", False, "222222", 11),
        ("4. Only once you've filled in F-H for a row, unhide columns I\u2013L (right-click the column headers \u2192 Unhide) to see what the engine actually said.", False, "222222", 11),
        ("5. \u201cMatch?\u201d (M) and \u201cRisk Match?\u201d (N) fill in automatically. If they disagree, pick why in \u201cDisagreement Type\u201d (O) and note what should change in \u201cFollow-up Notes\u201d (P).", False, "222222", 11),
        ("6. Once you've done this for 10-15 posts, check the \u201cSummary\u201d tab for your overall match rate and what kind of mistakes are most common.", False, "222222", 11),
        ("Why the engine columns start hidden", True, NAVY, 13),
        ("If you see the engine's answer before forming your own, you'll unconsciously anchor to it -- and the whole point of this exercise is finding out whether the engine's judgment "
         "actually matches yours when you're not looking at it first. It's slightly more friction, on purpose.", False, SLATE, 11),
    ]
    r = 4
    for text, bold, color, size in rows:
        _body(ws1, f"A{r}", text, bold=bold, size=size, color=color)
        ws1.row_dimensions[r].height = 22 if bold else 40
        r += 1

    # ================= Sheet 2: Validation Log =================
    ws2 = wb.create_sheet("Validation Log")
    ws2.sheet_view.showGridLines = False

    headers = [
        ("#", 5, NAVY), ("Post ID", 12, NAVY), ("Influencer\nHandle", 16, NAVY),
        ("Post URL", 26, NAVY), ("Timestamp", 14, NAVY),
        ("Your\nVerdict", 15, "2E7D32"), ("Your Risk\nGuess", 12, "2E7D32"), ("Your\nReasoning", 30, "2E7D32"),
        ("Engine\nVerdict", 15, SLATE), ("Engine\nRisk Level", 12, SLATE), ("Engine\nFlags", 24, SLATE), ("Engine Visual\nStatus", 14, SLATE),
        ("Match?", 11, GOLD), ("Risk\nMatch?", 11, GOLD),
        ("Disagreement Type", 26, GOLD), ("Follow-up Notes", 30, GOLD),
    ]
    for i, (text, width, fill) in enumerate(headers, start=1):
        _header_cell(ws2, 1, i, text, width, fill)
    ws2.row_dimensions[1].height = 34

    DATA_START = 2
    n = max(len(reports), 1)
    LAST_ROW = DATA_START + n + 20  # a little headroom for adding more posts later

    for idx, rep in enumerate(reports):
        r = DATA_START + idx
        ws2.cell(row=r, column=1, value=idx + 1)
        ws2.cell(row=r, column=2, value=get(rep, "post_id"))
        ws2.cell(row=r, column=3, value=get(rep, "influencer_handle"))
        ws2.cell(row=r, column=4, value=get(rep, "post_url"))
        ws2.cell(row=r, column=5, value=str(get(rep, "timestamp")))
        # F, G, H (Your Verdict / Risk / Reasoning) intentionally left blank
        ws2.cell(row=r, column=9, value=get(rep, "caption_status"))
        risk = get(rep, "risk_level") or "NONE"
        ws2.cell(row=r, column=10, value=risk)
        flags = get(rep, "caption_flags", [])
        ws2.cell(row=r, column=11, value=", ".join(flags) if flags else "")
        ws2.cell(row=r, column=12, value=get(rep, "visual_status"))

    for r in range(DATA_START, LAST_ROW + 1):
        ws2.cell(row=r, column=13, value=f'=IF(F{r}="","",IF(F{r}=I{r},"MATCH","MISMATCH"))')
        ws2.cell(row=r, column=14, value=f'=IF(OR(F{r}="",G{r}=""),"",IF(G{r}=J{r},"MATCH","MISMATCH"))')
        for c in range(1, 17):
            cell = ws2.cell(row=r, column=c)
            cell.border = border_all
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(name=FONT_NAME, size=10)
        ws2.row_dimensions[r].height = 30

    # Data validation dropdowns
    dv_verdict = DataValidation(type="list", formula1=f'"{",".join(VERDICT_OPTIONS)}"', allow_blank=True)
    dv_risk = DataValidation(type="list", formula1=f'"{",".join(RISK_OPTIONS)}"', allow_blank=True)
    dv_disagreement = DataValidation(type="list", formula1=f'"{",".join(DISAGREEMENT_TYPES)}"', allow_blank=True)
    for dv in (dv_verdict, dv_risk, dv_disagreement):
        ws2.add_data_validation(dv)
    dv_verdict.add(f"F{DATA_START}:F{LAST_ROW}")
    dv_risk.add(f"G{DATA_START}:G{LAST_ROW}")
    dv_disagreement.add(f"O{DATA_START}:O{LAST_ROW}")

    # Conditional formatting on Match columns
    for col_letter in ("M", "N"):
        rng = f"{col_letter}{DATA_START}:{col_letter}{LAST_ROW}"
        ws2.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"MATCH"'], fill=PatternFill("solid", fgColor=GREEN)))
        ws2.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"MISMATCH"'], fill=PatternFill("solid", fgColor=RED)))

    ws2.freeze_panes = "F2"
    # Hide the engine-answer columns until the reviewer is ready (see How To Use, step 4)
    for col_letter in ("I", "J", "K", "L"):
        ws2.column_dimensions[col_letter].hidden = True

    # ================= Sheet 3: Summary =================
    ws3 = wb.create_sheet("Summary")
    ws3.sheet_view.showGridLines = False
    ws3.column_dimensions["A"].width = 34
    ws3.column_dimensions["B"].width = 14
    ws3.column_dimensions["C"].width = 14

    _title(ws3, "A1", "Validation Summary")
    ws3.merge_cells("A1:C1")
    _body(ws3, "A3", "Fills in automatically as you complete rows in the Validation Log tab.", italic=True, size=10, color=SLATE)
    ws3.merge_cells("A3:C3")

    _header_cell(ws3, 5, 1, "Metric", 34, NAVY)
    _header_cell(ws3, 5, 2, "Count", 14, NAVY)
    _header_cell(ws3, 5, 3, "% of Reviewed", 14, NAVY)

    metrics = [
        ("Posts Reviewed (by you)", f'=COUNTIF(\'Validation Log\'!F:F,"COMPLIANT")+COUNTIF(\'Validation Log\'!F:F,"FLAGGED")+COUNTIF(\'Validation Log\'!F:F,"NEEDS EXPERT REVIEW")+COUNTIF(\'Validation Log\'!F:F,"PENDING REVIEW")'),
        ("Matches", '=COUNTIF(\'Validation Log\'!M:M,"MATCH")'),
        ("Mismatches", '=COUNTIF(\'Validation Log\'!M:M,"MISMATCH")'),
    ]
    r = 6
    for label, formula in metrics:
        ws3.cell(row=r, column=1, value=label).font = Font(name=FONT_NAME, size=10)
        ws3.cell(row=r, column=2, value=formula).font = Font(name=FONT_NAME, size=10)
        for c in (1, 2, 3):
            ws3.cell(row=r, column=c).border = border_all
        r += 1
    ws3.cell(row=7, column=3, value='=IFERROR(B7/$B$6,"")').number_format = "0%"
    ws3.cell(row=8, column=3, value='=IFERROR(B8/$B$6,"")').number_format = "0%"
    ws3.cell(row=6, column=1).font = Font(name=FONT_NAME, bold=True, size=10)

    ws3.cell(row=10, column=1, value="Disagreements by Type").font = Font(name=FONT_NAME, bold=True, size=12, color=NAVY)
    ws3.merge_cells("A10:C10")
    _header_cell(ws3, 12, 1, "Type", 34, SLATE)
    _header_cell(ws3, 12, 2, "# Occurrences", 14, SLATE)
    r = 13
    for dtype in DISAGREEMENT_TYPES:
        if dtype == "N/A - Match":
            continue
        ws3.cell(row=r, column=1, value=dtype).font = Font(name=FONT_NAME, size=10)
        ws3.cell(row=r, column=2, value=f'=COUNTIF(\'Validation Log\'!O:O,"{dtype}")').font = Font(name=FONT_NAME, size=10)
        for c in (1, 2):
            ws3.cell(row=r, column=c).border = border_all
        r += 1

    wb.save(filename)
    print(f"Wrote {filename} with {len(reports)} posts pre-filled from pipeline output.")
    return filename


if __name__ == "__main__":
    sample = [
        {"post_id": "SAMPLE1", "influencer_handle": "@example_account", "post_url": "https://instagram.com/p/example1/",
         "timestamp": "2026-07-14T00:00:00+0000", "caption_status": "FLAGGED",
         "caption_flags": ["approved_label", "placement"], "visual_status": "NO_TEXT_DETECTED", "risk_level": "HIGH"},
        {"post_id": "SAMPLE2", "influencer_handle": "@example_account", "post_url": "https://instagram.com/p/example2/",
         "timestamp": "2026-07-14T00:00:00+0000", "caption_status": "COMPLIANT",
         "caption_flags": [], "visual_status": "DISCLOSURE_VISIBLE", "risk_level": None},
    ]
    write_validation_log(sample, "data/validation_log_SAMPLE.xlsx")
